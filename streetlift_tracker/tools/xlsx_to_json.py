# -*- coding: utf-8 -*-
"""Extrait Programme_Streetlifting_v3-3.xlsx vers assets/programme_v33.json.
Traduit les formules Excel en specs recalculables par l'app (fidèle au Pilotage)."""
import argparse, json, re, sys, datetime
from pathlib import Path
from program_math import rest_seconds, round_step
from openpyxl import load_workbook

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('source', nargs='?', default='Programme_Streetlifting_v3-3.xlsx')
parser.add_argument('--output', type=Path, default=Path(__file__).resolve().parents[1] / 'assets/programme_v33.json')
args = parser.parse_args()
SRC = args.source
wbF = load_workbook(SRC)                    # formules
wbV = load_workbook(SRC, data_only=True)    # valeurs recalculées

BLOCKS = [
    ("P0", "Phase 0 (S1-3) Tests", "Phase 0 — Tests", "808080"),
    ("B1", "Bloc 1 (S4-11) Hypertrophie", "Bloc 1 — Hypertrophie", "2E6E8E"),
    ("B2", "Bloc 2 (S12-19) Force", "Bloc 2 — Force", "1F4E79"),
    ("B3", "Bloc 3 (S20-25) Force max", "Bloc 3 — Force max", "7A1F2B"),
    ("B4", "Bloc 4 (S26-31) Endurance", "Bloc 4 — Endurance", "2E7D32"),
    ("B5", "Bloc 5 (S32-40) Peaking", "Bloc 5 — Peaking", "B8860B"),
]

RE_BAN = re.compile(r"^S(\d+) · (.+?) · J(\d) — (.+?)(?: · (.+))?$")
RE_BAN7 = re.compile(r"^S(\d+) · J7 — (.+)$")
RE_SYS = re.compile(r"^=MAX\(0,ROUND\(\(\(Pilotage!\$B\$4\+Pilotage!\$B\$(\d+)\)\*([\d.]+)-Pilotage!\$B\$4\)/2\.5,0\)\*2\.5\)$")
RE_BAR = re.compile(r"^=ROUND\(Pilotage!\$B\$11\*([\d.]+)/2\.5,0\)\*2\.5$")
RE_ACC = re.compile(r"^=ROUND\(Pilotage!\$B\$(\d+)\*\(1\+\(Pilotage!\$D\$\d+\+2\)/30\)/\(1\+\((\d+)\+2\)/30\)/([\d.]+),0\)\*([\d.]+)$")
RE_VOL = re.compile(r'^="(.*?)"&ROUND\(([\d.]+)\*Pilotage!\$B\$(\d+)/(\d+),0\)&"(.*?)"$')

def load_spec(cell_f, cell_v):
    f = cell_f.value
    if f is None:
        return {"type": "none"}
    if isinstance(f, (int, float)):
        return {"type": "fixed", "kg": float(f)}
    if isinstance(f, str) and f.startswith("="):
        m = RE_SYS.match(f)
        if m: return {"type": "system", "ref": f"B{m.group(1)}", "pct": float(m.group(2))}
        m = RE_BAR.match(f)
        if m: return {"type": "barbell", "ref": "B11", "pct": float(m.group(1))}
        m = RE_ACC.match(f)
        if m: return {"type": "acc", "ref": f"B{m.group(1)}", "dayReps": int(m.group(2)), "step": float(m.group(3))}
        raise ValueError(f"Formule de charge non traduite {cell_f.coordinate}: {f}")
    return {"type": "fixed", "kg": 0.0}

def sets_spec(cell_f, cell_v):
    f = cell_f.value
    if isinstance(f, str) and f.startswith("="):
        m = RE_VOL.match(f)
        if m:
            return {"type": "volume", "prefix": m.group(1), "coef": float(m.group(2)),
                    "ref": f"B{m.group(3)}", "div": int(m.group(4)), "suffix": m.group(5)}
        raise ValueError(f"Formule de volume non traduite {cell_f.coordinate}: {f}")
    return {"type": "text", "value": str(f or "")}

def fill_rgb(cell):
    try:
        if cell.fill and cell.fill.fill_type == "solid":
            return str(cell.fill.start_color.rgb)
    except Exception:
        pass
    return None

weeks = {}
unparsed = []
for key, sheet, blabel, bcolor in BLOCKS:
    wsF, wsV = wbF[sheet], wbV[sheet]
    cur_day = None
    for r in range(4, wsF.max_row + 1):
        a = wsF.cell(row=r, column=1).value
        if isinstance(a, str):
            m = RE_BAN.match(a)
            if m:
                n, dates, j, title, cycle = int(m.group(1)), m.group(2), int(m.group(3)), m.group(4), m.group(5) or ""
                wk = weeks.setdefault(n, {"n": n, "dates": dates, "block": blabel,
                                          "blockKey": key, "color": bcolor, "days": []})
                cur_day = {"j": j, "title": title.strip(), "cycle": cycle.strip(),
                           "conduite": "", "exercises": []}
                wk["days"].append(cur_day)
                continue
            m = RE_BAN7.match(a)
            if m:
                n = int(m.group(1))
                wk = weeks.setdefault(n, {"n": n, "dates": "", "block": blabel,
                                          "blockKey": key, "color": bcolor, "days": []})
                wk["days"].append({"j": 7, "title": "REPOS COMPLET", "cycle": "",
                                   "conduite": m.group(2), "exercises": []})
                cur_day = None
                continue
            if a.startswith("CONDUITE") and cur_day is not None:
                cur_day["conduite"] = a
                continue
        # ligne d'exercice : col B = J1..J6
        b = wsF.cell(row=r, column=2).value
        if cur_day is not None and isinstance(b, str) and b.startswith("J"):
            name = wsF.cell(row=r, column=3).value
            if not name:
                continue
            cfill = fill_rgb(wsF.cell(row=r, column=3))
            ex = {
                "id": f"{key}-{r}",
                "name": str(name),
                "sets": sets_spec(wsF.cell(row=r, column=4), wsV.cell(row=r, column=4)),
                "intensity": str(wsF.cell(row=r, column=5).value or ""),
                "load": load_spec(wsF.cell(row=r, column=6), wsV.cell(row=r, column=6)),
                "rest": str(wsF.cell(row=r, column=7).value or ""),
                "restSec": rest_seconds(wsF.cell(row=r, column=7).value),
                "tempo": str(wsF.cell(row=r, column=8).value or ""),
                "cue": str(wsF.cell(row=r, column=9).value or ""),
                "main": cfill == "FFE6F2EC",
                "prevention": cfill == "FFEDE7F6",
            }
            if ex["load"].get("note"):
                unparsed.append((ex["id"], wsF.cell(row=r, column=6).value))
            cur_day["exercises"].append(ex)

# ---- Pilotage ----
pv = wbV["Pilotage"]
def cellv(addr):
    v = pv[addr].value
    return v

pilotage = {
    "bodyweight": float(cellv("B4")),
    "mainLifts": [],
    "repMax": [],
    "accessories": [],
}
for row, key in [(8, "pull"), (9, "dip"), (10, "mu"), (11, "squat")]:
    pilotage["mainLifts"].append({
        "ref": f"B{row}", "key": key, "name": str(cellv(f"A{row}")),
        "oneRm": float(cellv(f"B{row}")), "unit": str(cellv(f"C{row}")),
        "target": float(cellv(f"D{row}")),
    })
for row in range(16, 21):
    pilotage["repMax"].append({
        "ref": f"B{row}", "name": str(cellv(f"A{row}")),
        "max": float(cellv(f"B{row}")), "target": float(cellv(f"D{row}")),
    })
for row in range(25, 46):
    nm = cellv(f"A{row}")
    if nm:
        pilotage["accessories"].append({
            "ref": f"B{row}", "name": str(nm), "refLoad": float(cellv(f"B{row}")),
            "refReps": int(cellv(f"D{row}")), "note": str(cellv(f"E{row}") or ""),
        })

out = {
    "meta": {"version": "3.3", "anchorMonday": "2026-07-13",
             "generated": str(datetime.date.today()), "weeks": len(weeks)},
    "pilotage": pilotage,
    "weeks": [weeks[n] for n in sorted(weeks)],
}
args.output.parent.mkdir(parents=True, exist_ok=True)
with args.output.open('w', encoding='utf-8') as fh:
    json.dump(out, fh, ensure_ascii=False, separators=(",", ":"))

nex = sum(len(d["exercises"]) for w in out["weeks"] for d in w["days"])
print(f"Semaines: {len(weeks)} | Jours: {sum(len(w['days']) for w in out['weeks'])} | Exercices: {nex}")
print("Formules non traduites:", len(unparsed))
for u in unparsed[:5]:
    print("  ", u)

# ---- vérification croisée S8 J1 : specs recalculées vs valeurs Excel ----
def compute(spec, P):
    if spec["type"] == "fixed": return spec["kg"]
    if spec["type"] == "system":
        rm = P[spec["ref"]]; pdc = P["B4"]
        return max(0.0, round_step((pdc + rm) * spec["pct"] - pdc, 2.5))
    if spec["type"] == "barbell":
        return round_step(P["B11"] * spec["pct"], 2.5)
    if spec["type"] == "acc":
        ref, rr = P[spec["ref"]], P["D" + spec["ref"][1:]]
        raw = ref * (1 + (rr + 2) / 30) / (1 + (spec["dayReps"] + 2) / 30)
        return round_step(raw, spec["step"])
    return None

P = {"B4": pilotage["bodyweight"]}
for L in pilotage["mainLifts"]: P[L["ref"]] = L["oneRm"]
for a in pilotage["accessories"]:
    P[a["ref"]] = a["refLoad"]; P["D" + a["ref"][1:]] = a["refReps"]
for m in pilotage["repMax"]: P[m["ref"]] = m["max"]

wk8 = next(w for w in out["weeks"] if w["n"] == 8)
d1 = next(d for d in wk8["days"] if d["j"] == 1)
wsV = wbV["Bloc 1 (S4-11) Hypertrophie"]
print("\nS8 J1 — spec vs Excel :")
for ex in d1["exercises"]:
    row = int(ex["id"].split("-")[1])
    exc = wsV.cell(row=row, column=6).value
    got = compute(ex["load"], P)
    flag = "OK" if (exc is None or got is None or abs(float(exc) - got) < 0.01) else "ECART!"
    print(f"  {flag} {ex['name'][:32]:34s} spec={got} excel={exc}")
